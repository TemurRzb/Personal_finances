import openpyxl

def parse_sberbank_statement(file_path):
    """Считывает Excel-файл выписки Сбербанка и возвращает список операций."""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except FileNotFoundError:
        return [] # Если файла нет, возвращаем пустой список
        
    sheet = wb.active
    col_indices = {}
    header_row_idx = None

    # Ищем шапку таблицы
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        row_str = [str(cell).strip().lower() for cell in row if cell is not None]
        if any("дата проводки" in text for text in row_str):
            header_row_idx = row_idx
            for col_idx, cell_value in enumerate(row):
                if not cell_value: continue
                header = str(cell_value).strip().lower()
                if "дата проводки" in header: col_indices['date'] = col_idx
                elif "сумма по дебету" in header: col_indices['debit'] = col_idx
                elif "сумма по кредиту" in header: col_indices['credit'] = col_idx
                elif "назначение платежа" in header: col_indices['purpose'] = col_idx
            break 
            
    if not header_row_idx:
        return []

    operations = []
    # Собираем данные
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        date_val = row[col_indices.get('date')]
        if date_val is None:
            continue

        debit_val = row[col_indices.get('debit')] if 'debit' in col_indices else 0
        credit_val = row[col_indices.get('credit')] if 'credit' in col_indices else 0
        purpose_val = row[col_indices.get('purpose')] if 'purpose' in col_indices else ""

        operations.append({
            'Дата': date_val,
            'Расход (Дебет)': float(debit_val if debit_val is not None else 0),
            'Приход (Кредит)': float(credit_val if credit_val is not None else 0),
            'Назначение': str(purpose_val).strip() if purpose_val else ""
        })

    return operations